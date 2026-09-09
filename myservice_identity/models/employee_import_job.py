import base64
import json
import logging
import threading
import time

import openpyxl

from odoo import _, api, fields, models, SUPERUSER_ID
from odoo.modules.registry import Registry

_logger = logging.getLogger(__name__)


# Maps xlsx header (normalised: stripped + lowercased) → key expected by
# nrs.sap.sync.log.process_sap_payload(). Same mapping the dev seed script
# (scripts/load_employees_xlsx.py) uses, so xlsx loaded here behaves
# identically whether it came in via API or via odoo shell.
COLUMN_MAP = {
    'ir no':                       'ir_number',
    'ir number':                   'ir_number',
    'personnel number':            'displayName',
    'display name':                'displayName',
    'name':                        'displayName',
    'rank':                        'jobTitle',
    'job title':                   'jobTitle',
    'phone no.':                   'phone',
    'phone':                       'phone',
    'email':                       'email',
    'employment status':           'employmentStatus',
    'present location':            'location',
    'group/administration office': 'department',
    'department':                  'department',
    'sco':                         'costCenter',
}

# Periodic commit interval — bounded so a mid-import crash never loses more
# than this many rows of work and the DB transaction stays small.
COMMIT_EVERY = 200

# Cap on per-job error retention. We keep the first N errors verbatim and
# only count the rest; otherwise an entirely bad sheet blows the field up.
MAX_ERROR_SAMPLES = 100

# What the status endpoint reports when the stored error list cannot be read
# back. `error_count` is a separate integer and stays true, so the operator is
# told the DETAIL is missing rather than being handed an empty list, which
# reads as "this job had no errors" — the exact confident-wrong-answer this
# handler used to produce.
UNREADABLE_ERROR_SAMPLE = {
    'row': 0,
    'ir_number': '',
    'message': 'Error detail for this job could not be read back — the stored '
               'list is not valid JSON. The error COUNT is still accurate; the '
               'rows themselves are in the Odoo log.',
}


def _norm_cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


class NrsEmployeeImportJob(models.Model):
    _name = 'nrs.employee.import.job'
    _description = 'Employee Bulk Import Job (xlsx)'
    _order = 'create_date desc'
    _rec_name = 'display_name'

    display_name = fields.Char(compute='_compute_display_name', store=True)
    state = fields.Selection([
        ('pending',    'Pending'),
        ('processing', 'Processing'),
        ('completed',  'Completed'),
        ('failed',     'Failed'),
    ], default='pending', required=True, index=True)

    file_data = fields.Binary(string='Source File', attachment=True, required=True)
    file_name = fields.Char(string='File Name', required=True)
    sheet_name = fields.Char(
        string='Sheet',
        help='Sheet to read. If blank, the first sheet is used.',
    )

    total_rows     = fields.Integer(string='Rows With IR')
    processed      = fields.Integer()
    created_count  = fields.Integer(string='Created')
    updated_count  = fields.Integer(string='Updated')
    skipped_count  = fields.Integer(string='Skipped')
    error_count    = fields.Integer(string='Errors')

    progress_pct = fields.Float(compute='_compute_progress_pct', string='Progress %')

    errors = fields.Text(
        help='JSON list of up to %d error rows {row, ir_number, message}.'
        % MAX_ERROR_SAMPLES,
    )
    fatal_message = fields.Text(
        help='Set on failed jobs that aborted before/during row processing.',
    )

    started_at  = fields.Datetime(readonly=True)
    finished_at = fields.Datetime(readonly=True)
    duration_ms = fields.Integer(readonly=True)

    @api.depends('file_name', 'create_date')
    def _compute_display_name(self):
        for job in self:
            stamp = (job.create_date or fields.Datetime.now()).strftime('%Y-%m-%d %H:%M')
            job.display_name = '%s — %s' % (stamp, job.file_name or 'employee-import')

    @api.depends('total_rows', 'processed')
    def _compute_progress_pct(self):
        for job in self:
            if job.total_rows:
                job.progress_pct = 100.0 * job.processed / job.total_rows
            else:
                job.progress_pct = 0.0

    # ------------------------------------------------------------------
    # Public API — invoked by myservice_api controller
    # ------------------------------------------------------------------

    def start_async(self):
        """Kick off processing in a daemon thread with a fresh Odoo env.

        We must NOT reuse self.env here — the request transaction will commit
        and close before the worker is done. The worker opens its own cursor
        via Registry(db_name).cursor() so it owns its lifecycle.
        """
        self.ensure_one()
        db_name = self.env.cr.dbname
        job_id = self.id

        # Commit the pending row so the worker thread can see it.
        self.env.cr.commit()

        thread = threading.Thread(
            target=self._worker_entrypoint,
            args=(db_name, job_id),
            name='nrs-employee-import-%d' % job_id,
            daemon=True,
        )
        thread.start()
        return job_id

    def to_status_dict(self):
        self.ensure_one()
        err_list = []
        if self.errors:
            # `errors` is written by `_run()` and by nothing else, always as
            # `json.dumps(<list of dicts>)`. So a value that will not parse
            # means a row hand-edited or written by an older format — a
            # ValueError (json.JSONDecodeError is one) and nothing else. The
            # narrow catch is the point: the broad one here could not tell bad
            # stored data from a typo in this method, and answered "no errors"
            # to both.
            try:
                err_list = json.loads(self.errors)
            except ValueError:
                _logger.warning(
                    'Employee import job %s: stored error list is not valid '
                    'JSON (%d chars) — reporting the gap, not "no errors"',
                    self.id, len(self.errors), exc_info=True,
                )
                err_list = [dict(UNREADABLE_ERROR_SAMPLE)]
            if not isinstance(err_list, list):
                # Valid JSON of the wrong shape: `err_list[:20]` below would
                # slice a string or raise on an int, so the shape is settled
                # here rather than at the point of use.
                _logger.warning(
                    'Employee import job %s: stored error list parsed as %s, '
                    'not a list — reporting the gap',
                    self.id, type(err_list).__name__,
                )
                err_list = [dict(UNREADABLE_ERROR_SAMPLE)]
        return {
            'id': self.id,
            'state': self.state,
            'file_name': self.file_name,
            'sheet_name': self.sheet_name or '',
            'total_rows': self.total_rows,
            'processed': self.processed,
            'created': self.created_count,
            'updated': self.updated_count,
            'skipped': self.skipped_count,
            'errors': self.error_count,
            'progress_pct': round(self.progress_pct, 2),
            'started_at': self.started_at and self.started_at.isoformat() or None,
            'finished_at': self.finished_at and self.finished_at.isoformat() or None,
            'duration_ms': self.duration_ms,
            'fatal_message': self.fatal_message or '',
            'error_samples': err_list[:20],
        }

    # ------------------------------------------------------------------
    # Worker
    # ------------------------------------------------------------------

    @staticmethod
    def _worker_entrypoint(db_name, job_id):
        """Top-level thread target. Owns its registry cursor + env."""
        try:
            with Registry(db_name).cursor() as cr:
                env = api.Environment(cr, SUPERUSER_ID, {})
                job = env['nrs.employee.import.job'].browse(job_id)
                job._run()
        except Exception:
            # catch-all: reported — this is a thread root, so anything at all
            # can arrive; the traceback is logged and the job is moved to
            # `failed` so the status endpoint stops saying "processing".
            _logger.exception('Employee import job %s crashed at thread root', job_id)
            # Best-effort: open a second cursor to mark the job failed so the
            # status endpoint shows something useful instead of "processing
            # forever".
            try:
                with Registry(db_name).cursor() as cr:
                    env = api.Environment(cr, SUPERUSER_ID, {})
                    env['nrs.employee.import.job'].browse(job_id).write({
                        'state': 'failed',
                        'fatal_message': 'Worker thread crashed — see odoo logs',
                        'finished_at': fields.Datetime.now(),
                    })
            except Exception:
                # catch-all: reported — by here the database itself is
                # unreachable, so the log line is the only channel left. What
                # it records is a real gap: the job stays at `processing`
                # forever and only this traceback says why.
                _logger.exception(
                    'Could not even mark import job %s as failed', job_id,
                )

    def _run(self):
        """Process the upload. Runs inside the worker's own transaction."""
        self.ensure_one()
        start = time.time()
        self.write({'state': 'processing', 'started_at': fields.Datetime.now()})
        self.env.cr.commit()

        try:
            rows = self._parse_xlsx()
        except Exception as e:
            # catch-all: reported — `_parse_xlsx` runs base64, zipfile and
            # openpyxl over a file an operator uploaded, and those raise a zoo
            # (BadZipFile, InvalidFileException, binascii.Error, KeyError,
            # IndexError on an empty workbook, MemoryError on a huge one) with
            # no shared base class worth naming. The job is marked `failed` and
            # the exception is carried to the operator verbatim.
            _logger.exception('Failed to parse xlsx for job %s', self.id)
            self.write({
                'state': 'failed',
                # The type as well as the text: several of the above stringify
                # to '' or to a bare quoted key, and 'Parse error: ' told the
                # operator nothing at all.
                'fatal_message': 'Parse error: %s: %s' % (type(e).__name__, e),
                'finished_at': fields.Datetime.now(),
                'duration_ms': int((time.time() - start) * 1000),
            })
            self.env.cr.commit()
            return

        self.write({'total_rows': len(rows)})
        self.env.cr.commit()

        SyncLog = self.env['nrs.sap.sync.log']
        Employee = self.env['hr.employee']
        created = updated = skipped = errored = 0
        sampled_errors = []

        for i, payload in enumerate(rows, 1):
            ir = payload.get('ir_number', '')
            try:
                # One savepoint per row, so a row that fails undoes ITSELF.
                # `self.env.cr.rollback()` used to stand in the handler below,
                # and a rollback is not row-scoped: it discards everything
                # since the last commit — up to COMMIT_EVERY - 1 employees
                # that imported perfectly — while `created` and `updated` are
                # plain Python integers that went on counting them. The job
                # then reported 200 created for a batch that wrote none.
                with self.env.cr.savepoint():
                    pre_existing = Employee.search(
                        [('ir_number', '=', ir)], limit=1)
                    _employee, log = SyncLog.process_sap_payload(
                        payload, sync_type='csv_import',
                    )
                    status = log.status
                    error_message = log.error_message
                if status == 'success':
                    if pre_existing:
                        updated += 1
                    else:
                        created += 1
                elif status == 'skipped':
                    skipped += 1
                else:
                    errored += 1
                    if len(sampled_errors) < MAX_ERROR_SAMPLES:
                        sampled_errors.append({
                            'row': i,
                            'ir_number': ir,
                            'message': error_message or 'unknown',
                        })
            except Exception as e:
                # catch-all: reported — a row is one call into
                # `process_sap_payload`, which reaches hr.employee, the SAP
                # field mapping and whatever a tenant has added to either, so
                # the set is open; the row is counted as an error and its
                # message is carried to the operator in `error_samples`.
                errored += 1
                if len(sampled_errors) < MAX_ERROR_SAMPLES:
                    # Logged only while we are still sampling, so a wholly bad
                    # 10k sheet cannot put 10k tracebacks in the log.
                    _logger.warning(
                        'Employee import job %s: row %s (IR %s) failed',
                        self.id, i, ir or '(blank)', exc_info=True,
                    )
                    sampled_errors.append({
                        'row': i,
                        'ir_number': ir,
                        # The type as well as the text — `str(KeyError('x'))`
                        # is "'x'" and several ORM errors stringify to ''.
                        'message': '%s: %s' % (type(e).__name__, e),
                    })

            if i % COMMIT_EVERY == 0:
                self.write({
                    'processed': i,
                    'created_count': created,
                    'updated_count': updated,
                    'skipped_count': skipped,
                    'error_count': errored,
                    'errors': json.dumps(sampled_errors),
                })
                self.env.cr.commit()

        # Final write
        self.write({
            'state': 'completed',
            'processed': len(rows),
            'created_count': created,
            'updated_count': updated,
            'skipped_count': skipped,
            'error_count': errored,
            'errors': json.dumps(sampled_errors),
            'finished_at': fields.Datetime.now(),
            'duration_ms': int((time.time() - start) * 1000),
        })
        self.env.cr.commit()

    # ------------------------------------------------------------------
    # Parsing
    # ------------------------------------------------------------------

    def _parse_xlsx(self):
        """Return a list of payload dicts ready for process_sap_payload.

        Decodes the stored base64 file and walks it with openpyxl in
        read-only mode (keeps memory flat for 10k+ row sheets).
        """
        self.ensure_one()
        raw = base64.b64decode(self.file_data)
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(raw), data_only=True, read_only=True)

        if self.sheet_name and self.sheet_name in wb.sheetnames:
            ws = wb[self.sheet_name]
        else:
            ws = wb[wb.sheetnames[0]]

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return []

        headers = [_norm_cell(h).lower() for h in header_row]

        out = []
        for row in rows_iter:
            if not any(c not in (None, '') for c in row):
                continue
            payload = {}
            for header, raw_cell in zip(headers, row):
                key = COLUMN_MAP.get(header)
                if not key:
                    continue
                payload[key] = _norm_cell(raw_cell)
            if payload.get('ir_number'):
                out.append(payload)
        return out
